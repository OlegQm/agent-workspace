"""Create bounded, read-only previews for supported workspace files."""

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from fastapi import HTTPException, status
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from PIL import Image

IMAGE_MEDIA_TYPES = {
    ".gif": "image/gif",
    ".jpeg": "image/jpeg",
    ".jpg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
IMAGE_FORMATS = {
    ".gif": "GIF",
    ".jpeg": "JPEG",
    ".jpg": "JPEG",
    ".png": "PNG",
    ".webp": "WEBP",
}
SUPPORTED_PREVIEW_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", *IMAGE_MEDIA_TYPES}
MAX_PREVIEW_BYTES = 50 * 1024 * 1024
MAX_TEXT_PREVIEW_BYTES = 2 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
MAX_EXCEL_ROWS = 2_000
MAX_EXCEL_COLUMNS = 100
MAX_STYLED_EXCEL_BYTES = 10 * 1024 * 1024
MAX_STYLED_EXCEL_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_BLOCKS = 5_000
MAX_ARCHIVE_FILES = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


def docx_length_pixels(value: Any) -> float | None:
    """Convert an optional Word EMU length to CSS pixels at 96 DPI."""
    return round(float(value) / 9_525, 2) if value is not None else None


def docx_color(value: Any) -> str | None:
    """Convert an optional Word RGB color to a CSS hexadecimal color."""
    return f"#{value}" if value is not None else None


def first_defined(*values: Any) -> Any:
    """Return the first value that is not None."""
    return next((value for value in values if value is not None), None)


def safe_docx_link(value: str) -> str | None:
    """Allow links that a browser can open without executing an active URL scheme."""
    if not value:
        return None
    parsed = urlparse(value)
    if parsed.scheme.lower() in {"http", "https", "mailto"}:
        return value
    if not parsed.scheme and value.startswith("#"):
        return value
    return None


def docx_highlight(value: Any) -> str | None:
    """Convert a Word text highlight enum to a readable CSS color."""
    if value is None:
        return None
    colors = {
        "BLACK": "#000000",
        "BLUE": "#0000FF",
        "BRIGHT_GREEN": "#00FF00",
        "DARK_BLUE": "#000080",
        "DARK_RED": "#800000",
        "DARK_YELLOW": "#808000",
        "GRAY_25": "#C0C0C0",
        "GRAY_50": "#808080",
        "GREEN": "#008000",
        "PINK": "#FF00FF",
        "RED": "#FF0000",
        "TEAL": "#008080",
        "TURQUOISE": "#00FFFF",
        "VIOLET": "#800080",
        "WHITE": "#FFFFFF",
        "YELLOW": "#FFFF00",
    }
    return colors.get(getattr(value, "name", str(value)))


def docx_run(run: Any, paragraph: Paragraph, href: str | None = None) -> dict[str, Any]:
    """Serialize one formatted DOCX text run without exposing document XML."""
    paragraph_font = paragraph.style.font if paragraph.style is not None else None
    character_font = run.style.font if run.style is not None else None
    size = first_defined(
        run.font.size,
        character_font.size if character_font is not None else None,
        paragraph_font.size if paragraph_font is not None else None,
    )
    color = first_defined(
        run.font.color.rgb,
        character_font.color.rgb if character_font is not None else None,
        paragraph_font.color.rgb if paragraph_font is not None else None,
    )
    return {
        "text": run.text,
        "href": href,
        "bold": bool(
            first_defined(
                run.bold,
                character_font.bold if character_font is not None else None,
                paragraph_font.bold if paragraph_font is not None else None,
            )
        ),
        "italic": bool(
            first_defined(
                run.italic,
                character_font.italic if character_font is not None else None,
                paragraph_font.italic if paragraph_font is not None else None,
            )
        ),
        "underline": bool(
            first_defined(
                run.underline,
                character_font.underline if character_font is not None else None,
                paragraph_font.underline if paragraph_font is not None else None,
            )
        ),
        "strike": bool(
            first_defined(
                run.font.strike,
                character_font.strike if character_font is not None else None,
                paragraph_font.strike if paragraph_font is not None else None,
            )
        ),
        "superscript": bool(
            first_defined(
                run.font.superscript,
                character_font.superscript if character_font is not None else None,
                paragraph_font.superscript if paragraph_font is not None else None,
            )
        ),
        "subscript": bool(
            first_defined(
                run.font.subscript,
                character_font.subscript if character_font is not None else None,
                paragraph_font.subscript if paragraph_font is not None else None,
            )
        ),
        "font_size": round(size.pt, 2) if size is not None else None,
        "font_color": docx_color(color),
        "font_name": first_defined(
            run.font.name,
            character_font.name if character_font is not None else None,
            paragraph_font.name if paragraph_font is not None else None,
        ),
        "highlight_color": docx_highlight(run.font.highlight_color),
    }


def docx_paragraph(paragraph: Paragraph) -> dict[str, Any]:
    """Serialize a DOCX paragraph with common layout and list information."""
    alignment_names = {0: "left", 1: "center", 2: "right", 3: "justify", 4: "justify"}
    paragraph_format = paragraph.paragraph_format
    style_name = paragraph.style.name if paragraph.style else None
    list_level: int | None = None
    list_kind: str | None = None
    # python-docx does not expose paragraph numbering through its public API.
    properties = paragraph._p.pPr
    if properties is not None and properties.numPr is not None:
        level = properties.numPr.ilvl
        list_level = int(level.val) if level is not None else 0
        list_kind = "number" if style_name and "number" in style_name.lower() else "bullet"
    runs: list[dict[str, Any]] = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            href = safe_docx_link(item.url)
            runs.extend(docx_run(run, paragraph, href) for run in item.runs)
        else:
            runs.append(docx_run(item, paragraph))
    line_spacing = paragraph_format.line_spacing
    line_spacing_pixels = docx_length_pixels(line_spacing) if hasattr(line_spacing, "emu") else None
    line_spacing_multiple = (
        round(float(line_spacing), 2)
        if isinstance(line_spacing, (int, float)) and not isinstance(line_spacing, bool)
        else None
    )
    return {
        "type": "paragraph",
        "text": paragraph.text,
        "style": style_name,
        "alignment": alignment_names.get(int(paragraph.alignment))
        if paragraph.alignment is not None
        else None,
        "list_kind": list_kind,
        "list_level": list_level,
        "space_before": docx_length_pixels(paragraph_format.space_before),
        "space_after": docx_length_pixels(paragraph_format.space_after),
        "left_indent": docx_length_pixels(paragraph_format.left_indent),
        "right_indent": docx_length_pixels(paragraph_format.right_indent),
        "first_line_indent": docx_length_pixels(paragraph_format.first_line_indent),
        "line_spacing_pixels": line_spacing_pixels,
        "line_spacing_multiple": line_spacing_multiple,
        "runs": runs,
    }


def docx_cell(cell: Any) -> dict[str, Any]:
    """Serialize one DOCX table cell with its paragraphs and visual properties."""
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    fill = shading.get(qn("w:fill")) if shading is not None else None
    vertical_names = {0: "top", 1: "center", 2: "bottom"}
    return {
        "paragraphs": [docx_paragraph(paragraph) for paragraph in cell.paragraphs],
        "fill_color": f"#{fill}" if fill and fill.lower() not in {"auto", "none"} else None,
        "vertical_alignment": (
            vertical_names.get(int(cell.vertical_alignment))
            if cell.vertical_alignment is not None
            else None
        ),
    }


def docx_table(table: Table) -> dict[str, Any]:
    """Serialize a DOCX table while retaining cell formatting and column widths."""
    header_row = bool(
        table.rows
        and any(
            run.bold
            for cell in table.rows[0].cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
        )
    )
    return {
        "type": "table",
        "rows": [[docx_cell(cell) for cell in row.cells] for row in table.rows],
        "column_widths": [docx_length_pixels(column.width) for column in table.columns],
        "header_row": header_row,
        "style": table.style.name if table.style is not None else None,
    }


def preview_kind(path: Path) -> str | None:
    """Return the supported preview kind for a file path."""
    suffix = path.suffix.lower()
    if suffix in IMAGE_MEDIA_TYPES:
        return "image"
    return suffix[1:] if suffix in SUPPORTED_PREVIEW_EXTENSIONS else None


def ensure_previewable(path: Path) -> str:
    """Validate a file before parsing it for an in-app preview."""
    kind = preview_kind(path)
    if kind is None:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Preview is not supported"
        )
    if not path.is_file():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Path is not a file")
    if path.stat().st_size > MAX_PREVIEW_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail="File is too large to preview",
        )
    if kind in {"docx", "xlsx"}:
        validate_office_archive(path)
    return kind


def validate_office_archive(path: Path) -> None:
    """Reject malformed or excessively expanded Office Open XML archives."""
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            uncompressed = sum(member.file_size for member in members)
            if len(members) > MAX_ARCHIVE_FILES or uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise HTTPException(
                    status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                    detail="Office document is too large to preview safely",
                )
    except BadZipFile as exc:
        raise HTTPException(status_code=422, detail="Invalid Office document") from exc


def serialize_cell(value: Any) -> Any:
    """Convert an Excel cell value into a JSON-safe scalar."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def archive_uncompressed_size(path: Path) -> int:
    """Return the total expanded size of an Office Open XML archive."""
    with ZipFile(path) as archive:
        return sum(member.file_size for member in archive.infolist())


def excel_theme_colors(workbook: Any) -> list[str]:
    """Extract the ordered color scheme from an XLSX theme definition."""
    if not workbook.loaded_theme:
        return []
    try:
        root = ElementTree.fromstring(workbook.loaded_theme)
    except ElementTree.ParseError:
        return []
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", namespace)
    if scheme is None:
        return []
    colors_by_name: dict[str, str] = {}
    for item in list(scheme):
        name = item.tag.rsplit("}", 1)[-1]
        color = next(iter(item), None)
        raw = color.get("val") if color is not None else None
        if (not raw or not re.fullmatch(r"[0-9A-Fa-f]{6,8}", raw)) and color is not None:
            raw = color.get("lastClr")
        colors_by_name[name] = (
            f"#{raw[-6:].upper()}" if raw and re.fullmatch(r"[0-9A-Fa-f]{6,8}", raw) else "#000000"
        )
    theme_index_order = (
        "lt1",
        "dk1",
        "lt2",
        "dk2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
    )
    return [colors_by_name.get(name, "#000000") for name in theme_index_order]


def apply_excel_tint(color: str, tint: float) -> str:
    """Apply an Excel tint to a six-digit hexadecimal color."""
    if not color.startswith("#") or len(color) != 7 or not tint:
        return color
    channels = [int(color[index : index + 2], 16) for index in (1, 3, 5)]
    adjusted = [
        round(channel * (1 + tint)) if tint < 0 else round(channel + (255 - channel) * tint)
        for channel in channels
    ]
    return "#" + "".join(f"{max(0, min(255, channel)):02X}" for channel in adjusted)


def excel_color(color: Any, theme_colors: list[str]) -> str | None:
    """Resolve an openpyxl color into a browser-safe hexadecimal color."""
    if color is None:
        return None
    resolved: str | None = None
    if (
        color.type == "rgb"
        and isinstance(color.rgb, str)
        and re.fullmatch(r"[0-9A-Fa-f]{6,8}", color.rgb)
    ):
        resolved = f"#{color.rgb[-6:].upper()}"
    elif color.type == "theme" and isinstance(color.theme, int) and color.theme < len(theme_colors):
        resolved = theme_colors[color.theme]
    elif (
        color.type == "indexed"
        and isinstance(color.indexed, int)
        and color.indexed < len(COLOR_INDEX)
    ):
        resolved = f"#{COLOR_INDEX[color.indexed][-6:].upper()}"
    return apply_excel_tint(resolved, float(color.tint or 0)) if resolved else None


def excel_border_side(side: Any, theme_colors: list[str]) -> dict | None:
    """Serialize one visible Excel border side."""
    if not side or not side.style:
        return None
    return {"style": side.style, "color": excel_color(side.color, theme_colors) or "#B7B7B7"}


def excel_cell_style(cell: Any, theme_colors: list[str]) -> dict:
    """Serialize the visual cell properties supported by the web preview."""
    font_color = excel_color(cell.font.color, theme_colors)
    fill_color = (
        excel_color(cell.fill.fgColor, theme_colors)
        if cell.fill.fill_type in {"solid", "darkGray", "mediumGray", "lightGray"}
        else None
    )
    alignment = cell.alignment
    border = cell.border
    return {
        "bold": bool(cell.font.bold),
        "italic": bool(cell.font.italic),
        "underline": bool(cell.font.underline),
        "font_size": float(cell.font.sz) if cell.font.sz else None,
        "font_color": font_color,
        "fill_color": fill_color,
        "horizontal": alignment.horizontal,
        "vertical": alignment.vertical,
        "wrap_text": bool(alignment.wrap_text),
        "indent": int(alignment.indent or 0),
        "border_left": excel_border_side(border.left, theme_colors),
        "border_right": excel_border_side(border.right, theme_colors),
        "border_top": excel_border_side(border.top, theme_colors),
        "border_bottom": excel_border_side(border.bottom, theme_colors),
    }


def format_excel_number(value: float, number_format: str) -> str:
    """Render common Excel numeric formats without exposing floating-point noise."""
    section = number_format.split(";", 1)[0]
    normalized = re.sub(r'"[^"]*"|\[[^]]*]|\\.', "", section)
    if "%" in normalized:
        decimal_match = re.search(r"[.,]([0#]+).*%", normalized)
        decimals = len(decimal_match.group(1)) if decimal_match else 0
        return f"{value * 100:.{decimals}f}%"
    decimal_match = re.search(r"\.([0#]+)", normalized)
    if decimal_match:
        decimals = len(decimal_match.group(1))
        use_grouping = "," in normalized.split(".", 1)[0]
        return f"{value:{',' if use_grouping else ''}.{decimals}f}"
    if re.search(r"[0#]", normalized) and normalized.lower() != "general":
        return f"{value:,.0f}" if "," in normalized else f"{value:.0f}"
    return format(value, ".15g")


def excel_cell_display(cell: Any) -> str:
    """Return the value as Excel would commonly display it to a reader."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M") if value.time() else value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return format_excel_number(value, cell.number_format or "General")
    return str(value)


def excel_column_width(width: float | None) -> int:
    """Convert an Excel character width into a bounded CSS pixel width."""
    return max(32, min(600, round((width or 8.43) * 7 + 5)))


def excel_row_height(height: float | None) -> int:
    """Convert an Excel point height into a bounded CSS pixel height."""
    return max(18, min(240, round((height or 15) * 4 / 3)))


def excel_preview(path: Path, requested_sheet: str | None = None) -> dict:
    """Read a bounded worksheet preview with safe Excel layout metadata."""
    styled = (
        path.stat().st_size <= MAX_STYLED_EXCEL_BYTES
        and archive_uncompressed_size(path) <= MAX_STYLED_EXCEL_UNCOMPRESSED_BYTES
    )
    workbook = load_workbook(path, read_only=not styled, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        if not sheet_names:
            return {
                "kind": "xlsx",
                "sheets": [],
                "sheet": None,
                "rows": [],
                "column_widths": [],
                "row_heights": [],
                "styles": [{}],
                "truncated": False,
                "styled": styled,
            }
        sheet_name = requested_sheet or sheet_names[0]
        if sheet_name not in sheet_names:
            raise HTTPException(status_code=404, detail="Worksheet not found")
        worksheet = workbook[sheet_name]
        row_count = min(worksheet.max_row, MAX_EXCEL_ROWS)
        column_count = min(worksheet.max_column, MAX_EXCEL_COLUMNS)
        truncated = worksheet.max_row > MAX_EXCEL_ROWS or worksheet.max_column > MAX_EXCEL_COLUMNS
        theme_colors = excel_theme_colors(workbook) if styled else []
        merged_anchors: dict[tuple[int, int], tuple[int, int]] = {}
        merged_covered: set[tuple[int, int]] = set()
        if styled:
            for merged_range in worksheet.merged_cells.ranges:
                if merged_range.min_row > row_count or merged_range.min_col > column_count:
                    continue
                max_row = min(merged_range.max_row, row_count)
                max_column = min(merged_range.max_col, column_count)
                merged_anchors[(merged_range.min_row, merged_range.min_col)] = (
                    max_row - merged_range.min_row + 1,
                    max_column - merged_range.min_col + 1,
                )
                for row_index in range(merged_range.min_row, max_row + 1):
                    for column_index in range(merged_range.min_col, max_column + 1):
                        if (row_index, column_index) != (
                            merged_range.min_row,
                            merged_range.min_col,
                        ):
                            merged_covered.add((row_index, column_index))

        styles: list[dict] = [{}]
        rows: list[list[dict | None]] = []
        for row_index, row in enumerate(
            worksheet.iter_rows(max_row=row_count, max_col=column_count), start=1
        ):
            serialized_row: list[dict | None] = []
            for column_index, cell in enumerate(row, start=1):
                if (row_index, column_index) in merged_covered:
                    serialized_row.append(None)
                    continue
                row_span, column_span = merged_anchors.get((row_index, column_index), (1, 1))
                style_id = int(cell.style_id) if styled else 0
                while len(styles) <= style_id:
                    styles.append({})
                if styled and not styles[style_id]:
                    styles[style_id] = excel_cell_style(cell, theme_colors)
                serialized_row.append(
                    {
                        "value": serialize_cell(cell.value),
                        "display": excel_cell_display(cell),
                        "style_id": style_id,
                        "row_span": row_span,
                        "column_span": column_span,
                    }
                )
            rows.append(serialized_row)

        column_widths = [
            excel_column_width(worksheet.column_dimensions[get_column_letter(index)].width)
            if styled
            else excel_column_width(None)
            for index in range(1, column_count + 1)
        ]
        row_heights = [
            excel_row_height(worksheet.row_dimensions[index].height)
            if styled
            else excel_row_height(None)
            for index in range(1, row_count + 1)
        ]
        return {
            "kind": "xlsx",
            "sheets": sheet_names,
            "sheet": sheet_name,
            "rows": rows,
            "column_widths": column_widths,
            "row_heights": row_heights,
            "styles": styles,
            "show_gridlines": worksheet.sheet_view.showGridLines is not False if styled else True,
            "truncated": truncated,
            "styled": styled,
        }
    finally:
        workbook.close()


def docx_preview(path: Path) -> dict:
    """Read common DOCX text formatting and layout in document order."""
    document = Document(path)
    blocks: list[dict] = []
    truncated = False
    for index, item in enumerate(document.iter_inner_content(), start=1):
        if index > MAX_DOCX_BLOCKS:
            truncated = True
            break
        if isinstance(item, Paragraph):
            blocks.append(docx_paragraph(item))
        elif isinstance(item, Table):
            blocks.append(docx_table(item))
    section = document.sections[0]
    return {
        "kind": "docx",
        "blocks": blocks,
        "truncated": truncated,
        "page": {
            "width": docx_length_pixels(section.page_width),
            "height": docx_length_pixels(section.page_height),
            "margin_top": docx_length_pixels(section.top_margin),
            "margin_right": docx_length_pixels(section.right_margin),
            "margin_bottom": docx_length_pixels(section.bottom_margin),
            "margin_left": docx_length_pixels(section.left_margin),
        },
    }


def pdf_preview(path: Path) -> dict:
    """Read PDF page metadata without returning the original document bytes."""
    with pymupdf.open(path) as document:
        return {
            "kind": "pdf",
            "page_count": document.page_count,
            "pages": [
                {"width": round(page.rect.width, 2), "height": round(page.rect.height, 2)}
                for page in document
            ],
        }


def text_preview(path: Path) -> dict:
    """Read a bounded UTF-8 text preview while preserving its whitespace."""
    with path.open("rb") as stream:
        content = stream.read(MAX_TEXT_PREVIEW_BYTES + 1)
    truncated = len(content) > MAX_TEXT_PREVIEW_BYTES
    return {
        "kind": "txt",
        "content": content[:MAX_TEXT_PREVIEW_BYTES].decode("utf-8-sig", errors="replace"),
        "truncated": truncated,
    }


def image_preview(path: Path) -> dict:
    """Validate a raster image and return display metadata without embedding its bytes."""
    suffix = path.suffix.lower()
    with Image.open(path) as image:
        width, height = image.size
        if image.format != IMAGE_FORMATS[suffix]:
            raise ValueError("Image content does not match its extension")
        if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
            raise HTTPException(
                status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                detail="Image dimensions are too large to preview safely",
            )
        image.verify()
    return {
        "kind": "image",
        "media_type": IMAGE_MEDIA_TYPES[suffix],
        "width": width,
        "height": height,
    }


def document_preview(path: Path, sheet: str | None = None) -> dict:
    """Return a structured preview for a supported file."""
    kind = ensure_previewable(path)
    try:
        if kind == "xlsx":
            return excel_preview(path, sheet)
        if kind == "docx":
            return docx_preview(path)
        if kind == "txt":
            return text_preview(path)
        if kind == "image":
            return image_preview(path)
        return pdf_preview(path)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The document could not be previewed") from exc


def render_pdf_page(path: Path, page_number: int) -> bytes:
    """Render one PDF page as PNG for a toolbar-free read-only preview."""
    ensure_previewable(path)
    try:
        with pymupdf.open(path) as document:
            if page_number < 1 or page_number > document.page_count:
                raise HTTPException(status_code=404, detail="PDF page not found")
            page = document.load_page(page_number - 1)
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            output = BytesIO(pixmap.tobytes("png"))
            return output.getvalue()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The PDF page could not be rendered") from exc
