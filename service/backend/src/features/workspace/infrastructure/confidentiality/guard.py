"""Prevent restricted workspace inputs from becoming user-downloadable outputs."""

from __future__ import annotations

import hashlib
import re
import urllib.parse
import zipfile
from collections import OrderedDict
from pathlib import Path
from threading import Lock
from typing import Any

import pymupdf
from docx import Document
from openpyxl import load_workbook

from features.workspace.domain.policies.paths import path_is_within
from features.workspace.infrastructure.confidentiality.models import (
    BlockedDerivative,
    FileStamp,
    RestrictedSource,
)
from features.workspace.infrastructure.confidentiality.normalization import (
    fixed_windows as _fixed_windows,
)
from features.workspace.infrastructure.confidentiality.normalization import shingles as _shingles
from features.workspace.infrastructure.confidentiality.normalization import words as _words

MAX_HASH_BYTES = 300 * 1024 * 1024
MAX_TEXT_BYTES = 25 * 1024 * 1024
MAX_SOURCE_FILES = 2_000
MIN_OVERLAP_WORDS = 80
EXCERPT_WORD_LIMIT = 25
SOURCE_COVERAGE_LIMIT = 0.60
CANDIDATE_COVERAGE_LIMIT = 0.85
PLAIN_TEXT_SUFFIXES = {
    ".csv",
    ".html",
    ".json",
    ".log",
    ".md",
    ".rtf",
    ".sql",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
MARKDOWN_LINK_RE = re.compile(r"\[([^\]]*)\]\(([^)]+)\)")
URL_RE = re.compile(r"https?://[^\s<>)]+", re.IGNORECASE)
DERIVATIVE_CACHE_ENTRIES = 256
_DERIVATIVE_CACHE: OrderedDict[tuple[str, str, str], bool] = OrderedDict()
_DERIVATIVE_CACHE_LOCK = Lock()
_DERIVATIVE_CHECK_LOCKS = tuple(Lock() for _ in range(64))


def _digest(path: Path) -> str | None:
    """Return a bounded SHA-256 digest for a regular file."""
    try:
        size = path.stat().st_size
    except OSError:
        return None
    if size > MAX_HASH_BYTES:
        return None
    hasher = hashlib.sha256()
    try:
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                hasher.update(chunk)
    except OSError:
        return None
    return hasher.hexdigest()


def _extract_plain_text(path: Path) -> str:
    """Read a bounded text-like file with replacement for invalid bytes."""
    with path.open("rb") as stream:
        return stream.read(MAX_TEXT_BYTES + 1)[:MAX_TEXT_BYTES].decode("utf-8", errors="replace")


def _extract_docx_text(path: Path) -> str:
    """Extract paragraph and table text from a DOCX document."""
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def _extract_xlsx_text(path: Path) -> str:
    """Extract displayed cell values from an XLSX workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    values: list[str] = []
    try:
        for sheet in workbook.worksheets:
            values.append(sheet.title)
            for row in sheet.iter_rows(values_only=True):
                values.extend(str(value) for value in row if value is not None)
    finally:
        workbook.close()
    return "\n".join(values)


def _extract_pdf_text(path: Path) -> str:
    """Extract bounded text from a PDF document."""
    values: list[str] = []
    size = 0
    with pymupdf.open(path) as document:
        for page in document:
            value = page.get_text("text")
            values.append(value)
            size += len(value.encode("utf-8", errors="ignore"))
            if size >= MAX_TEXT_BYTES:
                break
    return "\n".join(values)[:MAX_TEXT_BYTES]


def extract_text(path: Path) -> str:
    """Extract searchable text from supported business-document formats."""
    try:
        if path.stat().st_size > MAX_TEXT_BYTES:
            return ""
        suffix = path.suffix.casefold()
        if suffix in PLAIN_TEXT_SUFFIXES:
            return _extract_plain_text(path)
        if suffix == ".docx":
            return _extract_docx_text(path)
        if suffix == ".xlsx":
            return _extract_xlsx_text(path)
        if suffix == ".pdf":
            return _extract_pdf_text(path)
    except Exception:
        return ""
    return ""


def _regular_files(path: Path) -> list[Path]:
    """Return regular non-symlink files represented by a protected path."""
    if path.is_file() and not path.is_symlink():
        return [path]
    if not path.is_dir() or path.is_symlink():
        return []
    return [item for item in path.rglob("*") if item.is_file() and not item.is_symlink()]


def snapshot_workspace(workspace_root: Path) -> dict[str, FileStamp]:
    """Capture workspace file metadata without reading file contents."""
    snapshot: dict[str, FileStamp] = {}
    if not workspace_root.exists():
        return snapshot
    for path in workspace_root.rglob("*"):
        if not path.is_file() or path.is_symlink():
            continue
        try:
            stat = path.stat()
        except OSError:
            continue
        snapshot[path.relative_to(workspace_root).as_posix()] = FileStamp(
            size=stat.st_size,
            modified_ns=stat.st_mtime_ns,
        )
    return snapshot


def _snapshot_digest(snapshot: dict[str, FileStamp]) -> str:
    """Hash workspace metadata for bounded derivative-check cache keys."""
    hasher = hashlib.sha256()
    for path, stamp in sorted(snapshot.items()):
        hasher.update(path.encode("utf-8", errors="surrogatepass"))
        hasher.update(b"\0")
        hasher.update(str(stamp.size).encode())
        hasher.update(b"\0")
        hasher.update(str(stamp.modified_ns).encode())
        hasher.update(b"\n")
    return hasher.hexdigest()


class RestrictedContentGuard:
    """Detect restricted-content leakage in messages and generated artifacts."""

    def __init__(
        self,
        workspace_root: Path,
        protected: list[str],
        sources: tuple[RestrictedSource, ...],
    ) -> None:
        """Initialize the guard from trusted protected-path fingerprints."""
        self.workspace_root = workspace_root
        self.protected = tuple(protected)
        self.sources = sources
        self.source_digests = frozenset(
            source.digest for source in sources if source.digest is not None
        )

    @classmethod
    def build(cls, workspace_root: Path, protected: list[str]) -> RestrictedContentGuard:
        """Build ephemeral fingerprints for all currently protected files."""
        sources: list[RestrictedSource] = []
        for protected_path in protected:
            root = workspace_root / protected_path
            for path in _regular_files(root):
                if len(sources) >= MAX_SOURCE_FILES:
                    break
                text_words = _words(extract_text(path))
                sources.append(
                    RestrictedSource(
                        relative_path=path.relative_to(workspace_root).as_posix(),
                        digest=_digest(path),
                        words=text_words,
                        shingles=_shingles(text_words),
                        prohibited_excerpts=_fixed_windows(text_words, EXCERPT_WORD_LIMIT + 1),
                    )
                )
        return cls(workspace_root, protected, tuple(sources))

    @property
    def active(self) -> bool:
        """Return whether at least one restricted source is enforceable."""
        return bool(self.sources)

    def is_explicitly_protected(self, relative_path: str) -> bool:
        """Return whether a path is an original protected input."""
        return any(path_is_within(relative_path, item) for item in self.protected)

    def _substantial_overlap(self, words: tuple[str, ...]) -> bool:
        """Detect long verbatim or near-verbatim reproduction of a source."""
        if len(words) < MIN_OVERLAP_WORDS:
            return False
        candidate = _shingles(words)
        if not candidate:
            return False
        for source in self.sources:
            if len(source.words) < MIN_OVERLAP_WORDS or not source.shingles:
                continue
            common = len(candidate.intersection(source.shingles))
            source_coverage = common / len(source.shingles)
            candidate_coverage = common / len(candidate)
            if (
                source_coverage >= SOURCE_COVERAGE_LIMIT
                or candidate_coverage >= CANDIDATE_COVERAGE_LIMIT
            ):
                return True
        return False

    def _exceeds_excerpt_limit(self, words: tuple[str, ...]) -> bool:
        """Detect a contiguous source excerpt longer than the allowed maximum."""
        candidate_windows = _fixed_windows(words, EXCERPT_WORD_LIMIT + 1)
        if not candidate_windows:
            return False
        return any(
            bool(candidate_windows.intersection(source.prohibited_excerpts))
            for source in self.sources
            if source.prohibited_excerpts
        )

    def _archive_contains_source(self, path: Path) -> bool:
        """Detect an exact restricted file embedded in a generated ZIP archive."""
        if not self.source_digests or not zipfile.is_zipfile(path):
            return False
        try:
            with zipfile.ZipFile(path) as archive:
                for info in archive.infolist():
                    if info.is_dir() or info.file_size > MAX_HASH_BYTES:
                        continue
                    hasher = hashlib.sha256()
                    with archive.open(info) as stream:
                        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                            hasher.update(chunk)
                    if hasher.hexdigest() in self.source_digests:
                        return True
        except (OSError, ValueError, zipfile.BadZipFile):
            return False
        return False

    def derivative_reason(self, path: Path, digest: str | None = None) -> str | None:
        """Return why a generated file exposes restricted content, if it does."""
        digest = digest if digest is not None else _digest(path)
        if digest is not None and digest in self.source_digests:
            return "exact restricted copy"
        if self._archive_contains_source(path):
            return "archive containing a restricted source"
        words = _words(extract_text(path))
        if self._exceeds_excerpt_limit(words):
            return "restricted excerpt longer than 25 words"
        if self._substantial_overlap(words):
            return "substantial restricted-content reproduction"
        return None

    def _matches_unrestricted_file(
        self,
        relative_path: str,
        digest: str | None,
        size: int,
        candidates: dict[str, FileStamp] | None = None,
    ) -> bool:
        """Return whether content is an exact copy of an allowed existing file."""
        if digest is None:
            return False
        paths = candidates if candidates is not None else snapshot_workspace(self.workspace_root)
        for other_path, stamp in paths.items():
            if (
                other_path == relative_path
                or stamp.size != size
                or self.is_explicitly_protected(other_path)
            ):
                continue
            source = self.workspace_root / other_path
            if not source.is_file() or source.is_symlink():
                continue
            if candidates is not None:
                try:
                    current = source.stat()
                except OSError:
                    continue
                if current.st_size != stamp.size or current.st_mtime_ns != stamp.modified_ns:
                    continue
            if _digest(source) == digest:
                return True
        return False

    def is_derivative_path(
        self,
        relative_path: str,
        candidates: dict[str, FileStamp] | None = None,
    ) -> bool:
        """Check an existing non-source workspace file for restricted derivation."""
        if not self.active or self.is_explicitly_protected(relative_path):
            return False
        path = self.workspace_root / relative_path
        if not path.is_file() or path.is_symlink():
            return False
        digest = _digest(path)
        if self._matches_unrestricted_file(
            relative_path,
            digest,
            path.stat().st_size,
            candidates,
        ):
            return False
        return self.derivative_reason(path, digest) is not None

    def remove_generated_derivatives(self, before: dict[str, FileStamp]) -> list[BlockedDerivative]:
        """Delete newly created or modified files that reproduce restricted inputs."""
        blocked: list[BlockedDerivative] = []
        after = snapshot_workspace(self.workspace_root)
        for relative_path, stamp in after.items():
            if before.get(relative_path) == stamp or self.is_explicitly_protected(relative_path):
                continue
            path = self.workspace_root / relative_path
            digest = _digest(path)
            if self._matches_unrestricted_file(relative_path, digest, stamp.size, before):
                continue
            reason = self.derivative_reason(path, digest)
            if reason is None:
                continue
            try:
                path.unlink()
            except OSError:
                continue
            blocked.append(BlockedDerivative(relative_path=relative_path, reason=reason))
        return blocked

    def _references_sensitive_path(self, value: str, blocked_paths: tuple[str, ...] = ()) -> bool:
        """Return whether a URL or link destination exposes a sensitive path."""
        decoded = urllib.parse.unquote_plus(value).replace("\\", "/").casefold()
        paths = (*self.protected, *blocked_paths)
        for path in paths:
            normalized = path.casefold()
            name = Path(path).name.casefold()
            stem = Path(path).stem.casefold()
            if normalized in decoded or name in decoded:
                return True
            if len(stem) >= 8 and stem in decoded:
                return True
        return False

    def sanitize_text(self, value: str, blocked_paths: tuple[str, ...] = ()) -> tuple[str, bool]:
        """Remove download references and reject substantial restricted reproduction."""
        changed = False

        def replace_markdown(match: re.Match[str]) -> str:
            """Replace a Markdown link that targets restricted content."""
            nonlocal changed
            if not self._references_sensitive_path(match.group(2), blocked_paths):
                return match.group(0)
            changed = True
            return f"{match.group(1)} (restricted source; download unavailable)"

        sanitized = MARKDOWN_LINK_RE.sub(replace_markdown, value)

        def replace_url(match: re.Match[str]) -> str:
            """Replace a plain URL that targets restricted content."""
            nonlocal changed
            if not self._references_sensitive_path(match.group(0), blocked_paths):
                return match.group(0)
            changed = True
            return "[restricted download unavailable]"

        sanitized = URL_RE.sub(replace_url, sanitized)
        sanitized_words = _words(sanitized)
        if self._exceeds_excerpt_limit(sanitized_words) or self._substantial_overlap(
            sanitized_words
        ):
            return (
                (
                    "The response was blocked because it reproduced too much content from a "
                    "restricted analysis-only file. Request a narrower aggregate, conclusion, or "
                    "short excerpt instead."
                ),
                True,
            )
        return sanitized, changed

    def sanitize_payload(self, value: Any) -> Any:
        """Recursively sanitize user-visible strings in a runtime event payload."""
        if isinstance(value, str):
            return self.sanitize_text(value)[0]
        if isinstance(value, list):
            return [self.sanitize_payload(item) for item in value]
        if isinstance(value, dict):
            return {key: self.sanitize_payload(item) for key, item in value.items()}
        return value


def cached_derivative_path_check(
    workspace_root: Path,
    protected: list[str],
    relative_path: str,
) -> bool:
    """Cache a stable read-only derivative decision for an unchanged workspace."""
    if not protected:
        return False
    before = snapshot_workspace(workspace_root)
    signature = _snapshot_digest(before)
    policy = hashlib.sha256("\0".join(sorted(protected)).encode()).hexdigest()
    key = (str(workspace_root.resolve()), policy, f"{relative_path}\0{signature}")
    lock = _DERIVATIVE_CHECK_LOCKS[int(signature[:2], 16) % len(_DERIVATIVE_CHECK_LOCKS)]
    with lock:
        with _DERIVATIVE_CACHE_LOCK:
            cached = _DERIVATIVE_CACHE.get(key)
            if cached is not None:
                _DERIVATIVE_CACHE.move_to_end(key)
                return cached
        guard = RestrictedContentGuard.build(workspace_root, protected)
        result = guard.is_derivative_path(relative_path, before)
        if _snapshot_digest(snapshot_workspace(workspace_root)) != signature:
            return True
        with _DERIVATIVE_CACHE_LOCK:
            _DERIVATIVE_CACHE[key] = result
            _DERIVATIVE_CACHE.move_to_end(key)
            while len(_DERIVATIVE_CACHE) > DERIVATIVE_CACHE_ENTRIES:
                _DERIVATIVE_CACHE.popitem(last=False)
        return result
