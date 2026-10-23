"""Test bounded document preview generation."""

import base64
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from features.workspace.infrastructure.previews import rendering as previews
from features.workspace.infrastructure.previews.rendering import (
    apply_excel_tint,
    archive_uncompressed_size,
    document_preview,
    docx_color,
    docx_highlight,
    docx_length_pixels,
    docx_paragraph,
    docx_preview,
    ensure_previewable,
    excel_border_side,
    excel_cell_display,
    excel_color,
    excel_column_width,
    excel_row_height,
    excel_theme_colors,
    first_defined,
    format_excel_number,
    preview_kind,
    render_pdf_page,
    safe_docx_link,
    serialize_cell,
    validate_office_archive,
)


def add_test_hyperlink(paragraph, text: str, address: str) -> None:
    """Append a minimal external hyperlink to a test paragraph."""
    relationship_id = paragraph.part.relate_to(
        address,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def test_excel_preview_returns_values_and_sheet_names(tmp_path: Path) -> None:
    """Verify that XLSX files are returned as bounded structured values."""
    path = tmp_path / "report.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet.append(["Metric", "Value"])
    worksheet.append(["Revenue", 1812.42])
    worksheet.merge_cells("A1:B1")
    worksheet["A1"].font = Font(bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor="002060")
    worksheet["B2"].number_format = "#,##0.00"
    worksheet.column_dimensions["A"].width = 24
    worksheet.row_dimensions[1].height = 24
    workbook.save(path)

    preview = document_preview(path)

    assert preview["kind"] == "xlsx"
    assert preview["sheets"] == ["Summary"]
    assert preview["rows"][0][0]["column_span"] == 2
    title_style = preview["styles"][preview["rows"][0][0]["style_id"]]
    assert title_style["bold"] is True
    assert title_style["fill_color"] == "#002060"
    assert preview["rows"][0][1] is None
    assert preview["rows"][1][0]["value"] == "Revenue"
    assert preview["rows"][1][1]["display"] == "1,812.42"
    assert preview["column_widths"][0] > preview["column_widths"][1]
    assert preview["row_heights"][0] == 32


def test_docx_preview_preserves_paragraphs_and_tables(tmp_path: Path) -> None:
    """Verify that DOCX paragraphs and table values can be viewed read-only."""
    path = tmp_path / "brief.docx"
    document = Document()
    document.add_heading("Assistant", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Important").bold = True
    table = document.add_table(rows=1, cols=2)
    add_test_hyperlink(table.cell(0, 0).paragraphs[0], "Country", "https://example.com")
    table.cell(0, 1).text = "SK"
    document.save(path)

    preview = document_preview(path)

    assert preview["kind"] == "docx"
    assert preview["blocks"][0]["text"] == "Assistant"
    assert preview["blocks"][1]["runs"][0]["bold"] is True
    country = preview["blocks"][2]["rows"][0][0]["paragraphs"][0]
    assert country["text"] == "Country"
    assert country["runs"][0]["href"] == "https://example.com"
    assert preview["blocks"][2]["rows"][0][1]["paragraphs"][0]["text"] == "SK"
    assert preview["page"]["width"] > 700


def test_pdf_preview_and_page_rendering(tmp_path: Path) -> None:
    """Verify that PDF previews expose metadata and rendered PNG pages only."""
    path = tmp_path / "report.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Confidential report")
    document.save(path)
    document.close()

    preview = document_preview(path)
    image = render_pdf_page(path, 1)

    assert preview["kind"] == "pdf"
    assert preview["page_count"] == 1
    assert preview["pages"] == [{"width": 595.0, "height": 842.0}]
    assert image.startswith(b"\x89PNG")


def test_text_and_image_previews_are_bounded_and_structured(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Expose readable text and validated raster-image metadata without embedding image bytes."""
    text_path = tmp_path / "notes.txt"
    text_path.write_text("First line\nDruhý riadok\n", encoding="utf-8")
    image_path = tmp_path / "chart.png"
    image_path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
    )

    assert document_preview(text_path) == {
        "kind": "txt",
        "content": "First line\nDruhý riadok\n",
        "truncated": False,
    }
    assert document_preview(image_path) == {
        "kind": "image",
        "media_type": "image/png",
        "width": 1,
        "height": 1,
    }
    assert preview_kind(tmp_path / "PHOTO.JPEG") == "image"
    assert preview_kind(tmp_path / "animation.gif") == "image"
    assert preview_kind(tmp_path / "graphic.webp") == "image"

    monkeypatch.setattr(previews, "MAX_TEXT_PREVIEW_BYTES", 5)
    assert document_preview(text_path) == {
        "kind": "txt",
        "content": "First",
        "truncated": True,
    }


def test_image_preview_rejects_invalid_raster_content(tmp_path: Path) -> None:
    """Reject a file whose image extension does not match readable raster content."""
    path = tmp_path / "broken.png"
    path.write_bytes(b"not an image")

    with pytest.raises(HTTPException, match="could not be previewed") as error:
        document_preview(path)

    assert error.value.status_code == 422


def test_image_preview_rejects_mismatched_formats_and_excessive_dimensions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Reject misleading extensions and images that would expand beyond the pixel limit."""

    class FakeImage:
        """Provide controlled image metadata for validation branches."""

        def __init__(self, image_format: str, size: tuple[int, int]) -> None:
            """Store the synthetic image format and dimensions."""
            self.format = image_format
            self.size = size

        def __enter__(self):
            """Return this fake from its context manager."""
            return self

        def __exit__(self, *_: object) -> None:
            """Close the synthetic image without side effects."""

        def verify(self) -> None:
            """Represent a successful decoder verification."""

    path = tmp_path / "chart.png"
    path.write_bytes(b"placeholder")
    monkeypatch.setattr(previews.Image, "open", lambda _: FakeImage("JPEG", (10, 10)))
    with pytest.raises(HTTPException, match="could not be previewed"):
        document_preview(path)

    monkeypatch.setattr(previews.Image, "open", lambda _: FakeImage("PNG", (10_000, 10_000)))
    with pytest.raises(HTTPException, match="dimensions") as excessive:
        document_preview(path)
    assert excessive.value.status_code == 413


def test_preview_support_helpers_handle_invalid_inputs(
    tmp_path: Path,
) -> None:
    """Preserve preview format validation and numeric display helpers."""
    unsupported = tmp_path / "unsupported.bin"
    unsupported.write_bytes(b"data")
    assert preview_kind(unsupported) is None
    with pytest.raises(HTTPException, match="not supported"):
        ensure_previewable(unsupported)
    assert format_excel_number(0.125, "0.0%") == "12.5%"
    assert format_excel_number(1234.5, "#,##0.00") == "1,234.50"
    assert apply_excel_tint("#000000", 0.5) == "#808080"
    assert excel_column_width(1) == 32
    assert excel_column_width(1_000) == 600
    assert excel_row_height(1) == 18
    assert excel_row_height(1_000) == 240


def test_preview_rejects_bad_office_documents_and_pdf_page_ranges(tmp_path: Path) -> None:
    """Preserve client-facing errors for corrupt Office archives and invalid PDF pages."""
    bad_docx = tmp_path / "bad.docx"
    bad_docx.write_bytes(b"not a zip")
    with pytest.raises(HTTPException) as bad_archive:
        ensure_previewable(bad_docx)
    assert bad_archive.value.status_code == 422
    pdf = tmp_path / "one-page.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(pdf)
    document.close()
    with pytest.raises(HTTPException) as bad_page:
        render_pdf_page(pdf, 2)
    assert bad_page.value.status_code == 404


def test_preview_helper_contracts_cover_supported_serialization_and_safe_values(
    tmp_path: Path,
) -> None:
    """Preserve safe link, DOCX, and Excel helper transformations."""
    assert safe_docx_link("") is None
    assert safe_docx_link("https://example.test") == "https://example.test"
    assert safe_docx_link("mailto:user@example.invalid") == "mailto:user@example.invalid"
    assert safe_docx_link("#bookmark") == "#bookmark"
    assert safe_docx_link("javascript:alert(1)") is None
    assert docx_length_pixels(None) is None
    assert docx_length_pixels(9_525) == 1
    assert docx_color(None) is None
    assert docx_color("FF00FF") == "#FF00FF"
    assert docx_highlight(None) is None
    assert docx_highlight(SimpleNamespace(name="YELLOW")) == "#FFFF00"
    assert docx_highlight(SimpleNamespace(name="UNKNOWN")) is None
    assert first_defined(None, 0, 1) == 0
    assert serialize_cell(None) is None
    assert serialize_cell(True) is True
    assert serialize_cell(__import__("datetime").date(2026, 1, 2)) == "2026-01-02"
    assert serialize_cell(SimpleNamespace()) == "namespace()"
    assert apply_excel_tint("bad", 0.2) == "bad"
    assert apply_excel_tint("#010203", -0.5) == "#000102"
    assert format_excel_number(1234.5, "0") == "1234"
    assert format_excel_number(1234.5, "#,##0") == "1,234"
    assert format_excel_number(12.5, "General") == "12.5"

    rgb = SimpleNamespace(type="rgb", rgb="AA112233", theme=None, indexed=None, tint=0)
    themed = SimpleNamespace(type="theme", rgb=None, theme=0, indexed=None, tint=0.5)
    indexed = SimpleNamespace(type="indexed", rgb=None, theme=None, indexed=0, tint=0)
    invalid = SimpleNamespace(type="rgb", rgb="bad", theme=None, indexed=None, tint=0)
    assert excel_color(None, []) is None
    assert excel_color(rgb, []) == "#112233"
    assert excel_color(themed, ["#000000"]) == "#808080"
    assert excel_color(indexed, []) is not None
    assert excel_color(invalid, []) is None
    assert excel_border_side(None, []) is None
    assert excel_border_side(SimpleNamespace(style="thin", color=invalid), []) == {
        "style": "thin",
        "color": "#B7B7B7",
    }

    workbook = Workbook()
    cell = workbook.active["A1"]
    cell.value = __import__("datetime").datetime(2026, 1, 2, 3, 4)
    assert excel_cell_display(cell) == "02.01.2026 03:04"
    cell.value = __import__("datetime").date(2026, 1, 2)
    assert excel_cell_display(cell) == "02.01.2026"
    cell.value = False
    assert excel_cell_display(cell) == "FALSE"
    cell.value = None
    assert excel_cell_display(cell) == ""

    class ThemeWorkbook:
        """Carry a synthetic invalid theme for error-path characterization."""

        loaded_theme = b"not xml"

    assert excel_theme_colors(SimpleNamespace(loaded_theme=None)) == []
    assert excel_theme_colors(ThemeWorkbook()) == []
    assert preview_kind(tmp_path / "REPORT.PDF") == "pdf"


def test_preview_validation_and_parser_failure_contracts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Preserve preview rejections and error normalization for malformed inputs."""
    directory = tmp_path / "document.pdf"
    directory.mkdir()
    with pytest.raises(HTTPException, match="not a file"):
        ensure_previewable(directory)

    large = tmp_path / "large.pdf"
    large.write_bytes(b"12")
    monkeypatch.setattr(previews, "MAX_PREVIEW_BYTES", 1)
    with pytest.raises(HTTPException, match="too large"):
        ensure_previewable(large)
    monkeypatch.setattr(previews, "MAX_PREVIEW_BYTES", 50 * 1024 * 1024)

    archive = tmp_path / "large.docx"
    with ZipFile(archive, "w", ZIP_DEFLATED) as output:
        output.writestr("word/document.xml", b"x" * 3)
    monkeypatch.setattr(previews, "MAX_ARCHIVE_UNCOMPRESSED_BYTES", 2)
    with pytest.raises(HTTPException, match="too large"):
        validate_office_archive(archive)
    monkeypatch.setattr(previews, "MAX_ARCHIVE_UNCOMPRESSED_BYTES", 200 * 1024 * 1024)
    assert archive_uncompressed_size(archive) == 3

    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"not a pdf")
    with pytest.raises(HTTPException, match="could not be previewed"):
        document_preview(pdf)
    with pytest.raises(HTTPException, match="could not be rendered"):
        render_pdf_page(pdf, 1)


def test_excel_preview_sheet_selection_and_unstyled_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Preserve worksheet selection failure and safe unstyled rendering for large workbooks."""
    path = tmp_path / "report.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet.append([1])
    workbook.save(path)
    with pytest.raises(HTTPException, match="Worksheet not found"):
        document_preview(path, "Missing")

    monkeypatch.setattr(previews, "MAX_STYLED_EXCEL_BYTES", 0)
    preview = document_preview(path)
    assert preview["styled"] is False
    assert preview["styles"] == [{}]


def test_docx_preview_truncates_when_configured_limit_is_reached(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Preserve an explicit bounded DOCX traversal limit."""
    path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("one")
    document.save(path)
    monkeypatch.setattr(previews, "MAX_DOCX_BLOCKS", 0)
    preview = docx_preview(path)
    assert preview["truncated"] is True
    assert preview["blocks"] == []


def test_preview_remaining_workbook_and_document_traversal_branches(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Preserve empty sheets, skipped merges, missing theme schemes, and unknown DOCX blocks."""

    class EmptyWorkbook:
        """Provide the minimum workbook protocol for an empty-workbook response."""

        def __init__(self) -> None:
            """Create an empty sheet collection."""
            self.sheetnames: list[str] = []

        def close(self) -> None:
            """Satisfy the workbook cleanup contract."""

    unused = tmp_path / "unused.xlsx"
    unused.write_bytes(b"unused")
    monkeypatch.setattr(previews, "load_workbook", lambda *args, **kwargs: EmptyWorkbook())
    monkeypatch.setattr(previews, "archive_uncompressed_size", lambda _: 0)
    assert previews.excel_preview(unused)["sheets"] == []

    class NoSchemeWorkbook:
        """Carry valid XML without a color-scheme element."""

        loaded_theme = (
            b'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" />'
        )

    assert excel_theme_colors(NoSchemeWorkbook()) == []

    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "visible"
    worksheet.merge_cells("A10:B10")
    workbook.save(path)
    monkeypatch.undo()
    monkeypatch.setattr(previews, "MAX_EXCEL_ROWS", 1)
    assert document_preview(path)["rows"][0][0]["value"] == "visible"

    class UnknownDocument:
        """Expose an unsupported document block while preserving page metadata access."""

        def __init__(self) -> None:
            """Create one metadata-only document section."""
            self.sections = [
                SimpleNamespace(
                    page_width=None,
                    page_height=None,
                    top_margin=None,
                    right_margin=None,
                    bottom_margin=None,
                    left_margin=None,
                )
            ]

        def iter_inner_content(self):
            """Yield a non-paragraph, non-table block."""
            return iter([object()])

    monkeypatch.setattr(previews, "Document", lambda _: UnknownDocument())
    assert docx_preview(tmp_path / "unknown.docx")["blocks"] == []


def test_docx_numbered_paragraph_preserves_numbering_metadata() -> None:
    """Preserve numbered-list level extraction from python-docx's XML-backed paragraph state."""
    document = Document()
    paragraph = document.add_paragraph("numbered", style="List Number")
    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_ilvl().val = 2
    preview = docx_paragraph(paragraph)
    assert preview["list_kind"] == "number"
    assert preview["list_level"] == 2
